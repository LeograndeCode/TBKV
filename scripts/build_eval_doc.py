#!/usr/bin/env python3
"""
Build the paper-ready Evaluation section as a Word (.docx) document.

Sections:
  1. Models and datasets
  2. Metrics
  3. Figures (with explanations)
  4. Result tables (with explanations)
  5. Summary of findings

All numbers are the validated n_items=10 evaluations on a single Quadro RTX 6000.
Run:  python scripts/build_eval_doc.py
"""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

REPO = Path(__file__).resolve().parents[1]
PLOTS = REPO / "results" / "plots"
OUT = REPO / "results" / "TBKV_Evaluation.docx"

ACCENT = RGBColor(0x1F, 0x4E, 0x79)


# ── helpers ───────────────────────────────────────────────────────────────────
def h(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = ACCENT
    return p


def para(doc, text, italic=False, size=11):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.italic = italic
    r.font.size = Pt(size)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)
    return p


def figure(doc, png, caption, width=6.2):
    doc.add_picture(str(png), width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(caption)
    r.italic = True
    r.font.size = Pt(9.5)


def table(doc, headers, rows, bold_first_col=False):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, htext in enumerate(headers):
        hdr[i].text = ""
        run = hdr[i].paragraphs[0].add_run(htext)
        run.bold = True
        run.font.size = Pt(10)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(val))
            run.font.size = Pt(10)
            if i == 0 and bold_first_col:
                run.bold = True
    return t


# ══════════════════════════════════════════════════════════════════════════════
def build():
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    title = doc.add_heading("Evaluation", level=0)
    for r in title.runs:
        r.font.color.rgb = ACCENT
    para(doc,
         "All results below are obtained on a single Quadro RTX 6000 GPU using the "
         "identical pretrained ViTDet-B and ViViT-B weights across every method. "
         "FLOPs are measured with in-code operation counters (no analytic estimates); "
         "latency and peak GPU memory are measured with CUDA events under an identical "
         "warm-up and reset protocol for all methods. Unless stated otherwise, numbers "
         "are from a 10-video development subset used for fast iteration; relative "
         "efficiency trends are the quantities of interest.", italic=True, size=10)

    # ── 1. Models and datasets ────────────────────────────────────────────────
    h(doc, "1. Models and Datasets", 1)

    h(doc, "1.1 Models", 2)
    bullet(doc, "ViTDet-B — a plain (non-hierarchical) Vision Transformer backbone "
                "(12 transformer blocks, 4 global-attention + 8 windowed) with a simple "
                "feature-pyramid neck and a Mask R-CNN detection head. Used for video "
                "object detection at 672x672 input.")
    bullet(doc, "ViViT-B — a factorised video Vision Transformer with a spatial "
                "transformer applied per frame followed by a temporal transformer. Used "
                "for video action recognition. Every block is a global-attention block.")

    para(doc, "Compared efficiency methods (all training-free, applied on top of the "
              "same frozen weights):")
    bullet(doc, "TBKV (ours) — content-adaptive cross-frame reuse: cosine-similarity "
                "matching selects background tokens whose cached key/value (and, with "
                "token-skip, whole-block output) are reused from a periodically refreshed "
                "keyframe, so only unmatched/foreground tokens are recomputed.")
    bullet(doc, "Eventful — temporal token gating: recomputes only the top-k tokens with "
                "the largest change from the previous frame; the rest reuse cached results.")
    bullet(doc, "STGT — spatio-temporal gating that keeps the top-k most salient tokens "
                "per frame.")
    bullet(doc, "MaskVD — builds a token mask from the previous frame's detections and "
                "skips background tokens, refreshing periodically.")

    h(doc, "1.2 Datasets", 2)
    bullet(doc, "ImageNet VID — large-scale video object-detection benchmark (30 object "
                "classes). Each item is a full video; frames are evaluated sequentially so "
                "that cross-frame reuse methods can exploit temporal redundancy. Input "
                "resized to a 672 px long edge.")
    bullet(doc, "Kinetics-400 — action-recognition benchmark (400 classes). Each clip is a "
                "full ~250-frame video classified into a single action label.")

    # ── 2. Metrics ────────────────────────────────────────────────────────────
    h(doc, "2. Metrics", 1)
    bullet(doc, "mAP@50 — mean Average Precision at IoU 0.5 (ImageNet VID). Primary "
                "detection-accuracy metric; higher is better.")
    bullet(doc, "Top-1 / Top-5 accuracy — fraction of clips whose correct action label is "
                "the top prediction / among the top five (Kinetics-400); higher is better.")
    bullet(doc, "GFLOPs / frame — billions of floating-point operations per frame, from the "
                "operation counters. The hardware-independent measure of compute; lower is "
                "better. For cross-frame methods this is the cost of a steady-state "
                "(non-keyframe) frame.")
    bullet(doc, "Latency / frame (ms) — measured wall-clock time per frame with CUDA-event "
                "timing after warm-up. Reflects real speed including any bookkeeping "
                "overhead a method adds; lower is better.")
    bullet(doc, "Peak memory (MB) — peak GPU memory allocated per frame, capturing the cost "
                "of the caches a method must keep; lower is better.")

    # ── 3. Figures ────────────────────────────────────────────────────────────
    h(doc, "3. Figures", 1)

    h(doc, "3.1 Accuracy vs. compute — ImageNet VID", 2)
    figure(doc, PLOTS / "fig_pareto_vid.png",
           "Figure 1. Detection accuracy (mAP@50) versus compute (GFLOPs/frame) on "
           "ImageNet VID. The shaded band marks the efficient frontier occupied by the "
           "token-gating baselines.")
    para(doc,
         "Each competitor curve is an accuracy/efficiency sweep; up-and-to-the-left is "
         "better. Eventful and STGT form the frontier, retaining ~84 mAP@50 at 60-70 "
         "GFLOPs, and MaskVD sits nearby at 88.8 GFLOPs. Every TBKV operating point "
         "(including the aggressive, +second-method, and all-blocks variants) lies to the "
         "lower-right of this frontier: TBKV spends far more compute (126-157 GFLOPs) while "
         "reaching lower accuracy (70-74 mAP@50). The gap is structural — TBKV reuses only "
         "the key/value projections of matched tokens, which is a small fraction of the "
         "per-block cost, so the dominant MLP and the eight windowed blocks keep running.")

    h(doc, "3.2 Deployment cost — ImageNet VID", 2)
    figure(doc, PLOTS / "fig_latency_memory_vid.png",
           "Figure 2. Wall-clock latency (left) and peak GPU memory (right) per frame on "
           "ImageNet VID, measured with an identical timing harness across all methods.",
           width=6.6)
    para(doc,
         "This figure tests whether TBKV's theoretical FLOP savings translate into a "
         "deployment advantage. They do not. TBKV is the slowest method overall (102 "
         "ms/frame) — 46% slower than the plain baseline and about 1.8x slower than every "
         "token-gating competitor — because its per-frame cosine matching and gather/scatter "
         "overhead costs more than the compute it removes. On memory, the token-gating "
         "methods that cache dense per-token state (Eventful 2118 MB, STGT 1242 MB) are "
         "expensive, but MaskVD (927 MB) already matches TBKV (929 MB) while being 1.8x "
         "faster and more accurate. There is no operating point where TBKV is the best "
         "choice on speed or memory.")

    h(doc, "3.3 Accuracy vs. compute — Kinetics-400", 2)
    figure(doc, PLOTS / "fig_pareto_vivit.png",
           "Figure 3. Action-recognition accuracy (Top-1) versus compute (GFLOPs/frame) on "
           "Kinetics-400 with ViViT-B.", width=5.6)
    para(doc,
         "On Kinetics-400 the picture is the same. TBKV halves per-frame compute versus the "
         "baseline (7.25 vs 14.74 GFLOPs) but loses ~20 Top-1 points. Eventful reaches the "
         "same 80% accuracy as the baseline at 2.70 GFLOPs — 2.7x cheaper than TBKV and with "
         "no accuracy loss — and, like TBKV, requires no fine-tuning. TBKV is therefore "
         "strictly dominated on this benchmark as well.")

    # ── 4. Tables ─────────────────────────────────────────────────────────────
    h(doc, "4. Result Tables", 1)

    h(doc, "Table 1. ImageNet VID — accuracy vs. compute (ViTDet-B, 672 px, n=10)", 2)
    table(doc,
          ["Method", "GFLOPs/frame", "mAP@50 (%)"],
          [["ViTDet-B (baseline)", "174.5", "84.1"],
           ["Eventful (k=512)", "60.5", "84.0"],
           ["STGT (k=512)", "68.5", "83.8"],
           ["MaskVD", "88.8", "83.1"],
           ["TBKV", "156.5", "72.0"],
           ["TBKV (aggressive)", "130.3", "74.4"],
           ["TBKV + 2nd method", "132.0", "70.5"],
           ["TBKV all-blocks", "125.7", "73.5"]],
          bold_first_col=True)
    para(doc,
         "The token-gating methods recover almost all of the baseline accuracy while cutting "
         "compute by 2.5-2.9x. TBKV's best variant (all-blocks, 125.7 GFLOPs) still uses "
         "1.4-2.1x more compute than the frontier methods and trails them by ~10 mAP@50 "
         "points at the 10-video sample size. Extending TBKV to all blocks or stacking a "
         "second gating method lowers its floor only modestly and never reaches the frontier.",
         size=10)

    h(doc, "Table 2. ImageNet VID — deployment cost (identical timing harness, n=10)", 2)
    table(doc,
          ["Method", "GFLOPs", "Latency (ms)", "Peak mem (MB)", "mAP@50 (%)"],
          [["ViTDet-B (baseline)", "174.5", "69.6", "760", "84.1"],
           ["Eventful (k=512)", "60.5", "58.3", "2118", "84.0"],
           ["STGT (k=512)", "68.5", "56.2", "1242", "83.8"],
           ["MaskVD", "88.8", "56.0", "927", "83.1"],
           ["TBKV", "156.5", "101.8", "929", "72.0"]],
          bold_first_col=True)
    para(doc,
         "FLOP reductions do not imply speedups: TBKV records the highest latency of all "
         "methods despite removing FLOPs, because matching overhead dominates. The one axis "
         "where TBKV is competitive — memory — is already occupied more effectively by MaskVD, "
         "which matches TBKV's memory while being faster and more accurate.", size=10)

    h(doc, "Table 3. Kinetics-400 — accuracy vs. compute (ViViT-B, per frame, n=10)", 2)
    table(doc,
          ["Method", "GFLOPs/frame", "Top-1 (%)", "Fine-tuning"],
          [["ViViT-B (baseline)", "14.74", "80.0", "no"],
           ["Eventful (top-k=24)", "2.70", "80.0", "no"],
           ["TBKV", "7.25", "60.0", "no"]],
          bold_first_col=True)
    para(doc,
         "Eventful attains the baseline accuracy at less than half of TBKV's compute and "
         "without fine-tuning, removing any training-free advantage TBKV might have claimed. "
         "TBKV's ~2x FLOP reduction comes at a large accuracy cost.", size=10)

    # ── 5. Summary ────────────────────────────────────────────────────────────
    h(doc, "5. Summary of Findings", 1)
    bullet(doc, "On both benchmarks, TBKV is Pareto-dominated on compute: existing "
                "token-gating methods achieve equal accuracy at a fraction of the FLOPs.")
    bullet(doc, "TBKV provides no wall-clock speedup — it is in fact the slowest method on "
                "ImageNet VID because its matching overhead outweighs the compute it saves.")
    bullet(doc, "TBKV provides no memory advantage that is not already met by MaskVD at "
                "lower latency and higher accuracy.")
    bullet(doc, "TBKV has no training-free edge: the strongest competitor is equally "
                "training-free while being cheaper and as accurate.")
    para(doc,
         "The root cause is structural: TBKV reuses only the cheapest part of each block "
         "(the key/value projections of matched tokens) in a subset of blocks, so the "
         "dominant MLP cost, the eight windowed blocks, and the detection head continue to "
         "run every frame, while the matching step adds overhead. In its current form the "
         "method is not competitive as an efficiency contribution.", italic=True, size=10)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print(f"Saved: {OUT}")


if __name__ == "__main__":
    build()
